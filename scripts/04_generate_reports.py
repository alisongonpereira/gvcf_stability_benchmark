#!/usr/bin/env python3
"""
Step 4 — Report Generation

Reads all metrics JSON files from benchmarks/03_metrics/ and produces:
  • benchmarks/04_reports/benchmark_report.html  (interactive Plotly charts)
  • benchmarks/04_reports/benchmark_data.xlsx    (raw + summary sheets)

Usage:
    python3 04_generate_reports.py \
        --benchmark-dir /path/to/benchmarks \
        --output-dir    /path/to/benchmarks/04_reports
"""

import argparse
import json
import math
import os
import re
import sys
from datetime import datetime
from pathlib import Path

# Optional — Excel support
# NOTE: import openpyxl under env -u LD_PRELOAD; lxml (pulled by openpyxl)
# is a C extension that segfaults when libjemalloc is preloaded via LD_PRELOAD.
try:
    import openpyxl
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side
    )
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ─── Statistics helpers (pure Python, no numpy/scipy needed) ──────────────────

def _mean(vals):
    return sum(vals) / len(vals) if vals else 0.0


def _linreg(x, y):
    """Return (slope, intercept) for simple linear regression."""
    n = len(x)
    if n < 2:
        return None, None
    mx, my = _mean(x), _mean(y)
    denom = sum((xi - mx) ** 2 for xi in x)
    if denom == 0:
        return None, None
    slope = sum((x[i] - mx) * (y[i] - my) for i in range(n)) / denom
    intercept = my - slope * mx
    return slope, intercept


def compute_r2(x_vals, y_vals):
    """Coefficient of determination for linear fit.  Returns None if undefined."""
    x = [v for v in x_vals]
    y = [v for v in y_vals]
    if len(x) < 2 or len(y) < 2:
        return None
    slope, intercept = _linreg(x, y)
    if slope is None:
        return None
    my = _mean(y)
    ss_tot = sum((yi - my) ** 2 for yi in y)
    if ss_tot == 0:
        return 1.0
    ss_res = sum((y[i] - (slope * x[i] + intercept)) ** 2 for i in range(len(x)))
    return max(0.0, min(1.0, round(1 - ss_res / ss_tot, 4)))


def find_inflections(x_vals, y_vals, rel_threshold=0.5):
    """Return dataset sizes where the per-interval slope changes significantly."""
    if len(x_vals) < 3:
        return []
    slopes = []
    for i in range(1, len(x_vals)):
        dx = x_vals[i] - x_vals[i - 1]
        dy = y_vals[i] - y_vals[i - 1]
        slopes.append(dy / dx if dx else 0.0)
    mean_abs = _mean([abs(s) for s in slopes]) or 1.0
    inflections = []
    for i in range(1, len(slopes)):
        change = abs(slopes[i] - slopes[i - 1])
        if change > rel_threshold * mean_abs:
            inflections.append({
                "dataset_size":  x_vals[i],
                "slope_before":  round(slopes[i - 1], 6),
                "slope_after":   round(slopes[i], 6),
                "slope_change":  round(change, 6),
            })
    return inflections


# ─── Data loading ─────────────────────────────────────────────────────────────


SOFTWARES  = ["glnexus", "parabricks", "gatk"]
SOFT_LABEL = {"glnexus": "GLnexus", "parabricks": "Parabricks", "gatk": "GATK"}
SOFT_COLOR = {"glnexus": "#1f77b4", "parabricks": "#ff7f0e", "gatk": "#2ca02c"}

METRICS = {
    "wall_time_s":          "Wall Time (s)",
    "cpu_pct_avg":          "CPU avg (%)",
    "cpu_pct_max":          "CPU peak (%)",
    "gpu_util_pct_avg":     "GPU util avg (%)",
    "gpu_util_pct_max":     "GPU util peak (%)",
    "gpu_mem_used_mb_max":  "GPU Mem peak (MB)",
    "ram_used_gb_avg":      "RAM avg (GB)",
    "ram_used_gb_max":      "RAM peak (GB)",
    "disk_used_gb_max":     "Disk peak (GB)",
}


def load_all_metrics(metrics_dir: Path) -> dict:
    """
    Returns dict[software][size] = metrics_dict, merging top-level
    wall_time_s with the nested resources dict.
    """
    data: dict[str, dict[int, dict]] = {s: {} for s in SOFTWARES}

    for path in sorted(metrics_dir.glob("metrics_*.json")):
        m = re.match(r"metrics_(\w+)_(\d+)\.json", path.name)
        if not m:
            continue
        software, size = m.group(1), int(m.group(2))
        if software not in data:
            data[software] = {}
        try:
            with open(path) as f:
                rec = json.load(f)
        except Exception as e:
            print(f"[WARN] Could not read {path}: {e}")
            continue

        # Flatten: top-level + resources dict
        flat = {
            "software":      rec.get("software", software),
            "dataset_size":  rec.get("dataset_size", size),
            "status":        rec.get("status", "unknown"),
            "exit_code":     rec.get("exit_code", -1),
            "output_valid":  rec.get("output_valid", False),
            "variant_count": rec.get("variant_count", 0),
            "wall_time_s":   rec.get("wall_time_s", 0),
            "start_time":    rec.get("start_time", ""),
            "end_time":      rec.get("end_time", ""),
        }
        for k, v in rec.get("resources", {}).items():
            flat[k] = v

        data[software][size] = flat

    return data


def load_variant_data(metrics_dir: Path) -> tuple[dict, dict]:
    """
    Load variant analysis outputs produced by 03_analyze_variants.py.

    Returns:
        analysis_data  – {software: {size: gq_data_dict}}
        overlap_data   – {size: overlap_dict}
    """
    analysis: dict[str, dict[int, dict]] = {s: {} for s in SOFTWARES}
    overlap:  dict[int, dict]            = {}

    for path in sorted(metrics_dir.glob("variant_analysis_*.json")):
        m = re.match(r"variant_analysis_(\w+)_(\d+)\.json", path.name)
        if not m:
            continue
        sw, size = m.group(1), int(m.group(2))
        try:
            with open(path) as f:
                analysis.setdefault(sw, {})[size] = json.load(f)
        except Exception as e:
            print(f"[WARN] Could not read {path}: {e}")

    for path in sorted(metrics_dir.glob("variant_overlap_*.json")):
        m = re.match(r"variant_overlap_(\d+)\.json", path.name)
        if not m:
            continue
        size = int(m.group(1))
        try:
            with open(path) as f:
                overlap[size] = json.load(f)
        except Exception as e:
            print(f"[WARN] Could not read {path}: {e}")

    return analysis, overlap


def build_series(data: dict, metric: str) -> dict[str, tuple[list, list]]:
    """
    Returns {software: (x_vals, y_vals)} for the given metric,
    sorted by dataset size, only for successful runs.
    """
    series: dict[str, tuple[list, list]] = {}
    for sw, size_dict in data.items():
        xs, ys = [], []
        for size in sorted(size_dict):
            rec = size_dict[size]
            if rec.get("status") != "success":
                continue
            val = rec.get(metric)
            if val is None or not isinstance(val, (int, float)):
                continue
            xs.append(size)
            ys.append(val)
        if xs:
            series[sw] = (xs, ys)
    return series


# ─── Variant analysis helpers ────────────────────────────────────────────────

def _gq_line_colors(base_hex: str, n_lines: int) -> list[str]:
    """Gradient of RGBA colours: light (α=0.25) → opaque (α=1.0)."""
    r = int(base_hex[1:3], 16)
    g = int(base_hex[3:5], 16)
    b = int(base_hex[5:7], 16)
    colors = []
    for i in range(n_lines):
        alpha = 0.25 + 0.75 * (i / max(n_lines - 1, 1))
        colors.append(f"rgba({r},{g},{b},{alpha:.2f})")
    return colors


def _gq_chart_js(
    div_id: str,
    sw: str,
    size_hist_list: list,   # sorted list of (size, hist_dict)
    title: str,
) -> str:
    """Return Plotly.newPlot(...) JS for a GQ density histogram."""
    bin_starts = list(range(0, 100, 5))
    bin_mids   = [b + 2.5 for b in bin_starts]
    base_color = SOFT_COLOR.get(sw, "#888")
    colors     = _gq_line_colors(base_color, len(size_hist_list))

    traces = []
    for i, (size, hist) in enumerate(size_hist_list):
        ys = [round(float(hist.get(str(b), 0)), 6) for b in bin_starts]
        traces.append(json.dumps({
            "x": bin_mids,
            "y": ys,
            "name": f"N={size}",
            "mode": "lines",
            "line": {"color": colors[i], "width": 2},
            "hovertemplate": f"N={size} | GQ=%{{x:.0f}} | density=%{{y:.4f}}<extra></extra>",
        }))

    layout = {
        "title": {"text": title, "font": {"size": 13}},
        "xaxis": {"title": "Genotype Quality (GQ)", "gridcolor": "#eee", "range": [0, 100]},
        "yaxis": {"title": "Density", "gridcolor": "#eee"},
        "margin": {"l": 60, "r": 10, "t": 40, "b": 60},
        "paper_bgcolor": "rgba(0,0,0,0)",
        "plot_bgcolor": "#fafafa",
        "font": {"family": "Segoe UI, Arial", "size": 11},
        "legend": {"orientation": "h", "y": -0.35, "font": {"size": 10}},
        "hovermode": "x unified",
    }
    return (
        f"Plotly.newPlot('{div_id}', [{', '.join(traces)}], "
        f"{json.dumps(layout)}, {{responsive: true}});"
    )


def _pct_str(count: int, total: int) -> str:
    if total == 0:
        return "0%"
    return f"{100 * count / total:.1f}%"


def _make_venn_svg(overlap: dict) -> str:
    """Return an inline SVG Venn diagram for 2 or 3 tools."""
    tools     = overlap.get("tools", [])
    total     = overlap.get("total_unique", 0)
    counts    = overlap.get("counts", {})
    exclusive = overlap.get("exclusive", {})
    all_common = overlap.get("all_common", 0)
    pairwise  = overlap.get("pairwise_only", {})
    size      = overlap.get("dataset_size", "?")

    def fmt(n):
        return f"{n:,}" if isinstance(n, int) else str(n)

    def label_block(cx, cy, main_val, sub_val):
        return (
            f'<text x="{cx}" y="{cy}" text-anchor="middle" '
            f'font-size="13" font-weight="bold">{fmt(main_val)}</text>'
            f'<text x="{cx}" y="{cy + 17}" text-anchor="middle" '
            f'font-size="11" fill="#666">{sub_val}</text>'
        )

    n_tools = len(tools)

    if n_tools == 2:
        t1, t2 = tools
        c1  = SOFT_COLOR.get(t1, "#888")
        c2  = SOFT_COLOR.get(t2, "#999")
        ex1 = exclusive.get(t1, 0)
        ex2 = exclusive.get(t2, 0)
        common = all_common

        svg = (
            '<svg viewBox="0 0 500 250" xmlns="http://www.w3.org/2000/svg" '
            'style="max-width:480px;width:100%">\n'
            # circles
            f'  <circle cx="175" cy="122" r="118" fill="{c1}" opacity="0.30" stroke="{c1}" stroke-width="2"/>\n'
            f'  <circle cx="325" cy="122" r="118" fill="{c2}" opacity="0.30" stroke="{c2}" stroke-width="2"/>\n'
            # tool names
            f'  <text x="112" y="17" text-anchor="middle" font-size="13" font-weight="bold" fill="{c1}">{SOFT_LABEL.get(t1, t1)}</text>\n'
            f'  <text x="388" y="17" text-anchor="middle" font-size="13" font-weight="bold" fill="{c2}">{SOFT_LABEL.get(t2, t2)}</text>\n'
            # counts
            + label_block(108, 118, ex1, _pct_str(ex1, total)) + "\n"
            + label_block(250, 118, common, _pct_str(common, total)) + "\n"
            + label_block(392, 118, ex2, _pct_str(ex2, total)) + "\n"
            # footer
            f'  <text x="250" y="244" text-anchor="middle" font-size="11" fill="#888">'
            f'Total unique: {fmt(total)}\u2002|\u2002N={size}</text>\n'
            "</svg>"
        )
        return svg

    elif n_tools >= 3:
        t1, t2, t3 = tools[0], tools[1], tools[2]
        c1  = SOFT_COLOR.get(t1, "#888")
        c2  = SOFT_COLOR.get(t2, "#999")
        c3  = SOFT_COLOR.get(t3, "#aaa")
        ex1 = exclusive.get(t1, 0)
        ex2 = exclusive.get(t2, 0)
        ex3 = exclusive.get(t3, 0)
        p12 = pairwise.get(f"{t1}_{t2}", 0)
        p13 = pairwise.get(f"{t1}_{t3}", 0)
        p23 = pairwise.get(f"{t2}_{t3}", 0)

        svg = (
            '<svg viewBox="0 0 500 400" xmlns="http://www.w3.org/2000/svg" '
            'style="max-width:480px;width:100%">\n'
            # circles (equilateral triangle layout)
            f'  <circle cx="175" cy="148" r="115" fill="{c1}" opacity="0.28" stroke="{c1}" stroke-width="2"/>\n'
            f'  <circle cx="325" cy="148" r="115" fill="{c2}" opacity="0.28" stroke="{c2}" stroke-width="2"/>\n'
            f'  <circle cx="250" cy="275" r="115" fill="{c3}" opacity="0.28" stroke="{c3}" stroke-width="2"/>\n'
            # tool names
            f'  <text x="95"  y="37" text-anchor="middle" font-size="13" font-weight="bold" fill="{c1}">{SOFT_LABEL.get(t1, t1)}</text>\n'
            f'  <text x="405" y="37" text-anchor="middle" font-size="13" font-weight="bold" fill="{c2}">{SOFT_LABEL.get(t2, t2)}</text>\n'
            f'  <text x="250" y="390" text-anchor="middle" font-size="13" font-weight="bold" fill="{c3}">{SOFT_LABEL.get(t3, t3)}</text>\n'
            # exclusive regions
            + label_block(90,  135, ex1, _pct_str(ex1, total)) + "\n"
            + label_block(410, 135, ex2, _pct_str(ex2, total)) + "\n"
            + label_block(250, 350, ex3, _pct_str(ex3, total)) + "\n"
            # pairwise-only regions
            + label_block(250, 108, p12, _pct_str(p12, total)) + "\n"
            + label_block(148, 248, p13, _pct_str(p13, total)) + "\n"
            + label_block(352, 248, p23, _pct_str(p23, total)) + "\n"
            # all-common
            + label_block(250, 203, all_common, _pct_str(all_common, total)) + "\n"
            # footer
            f'  <text x="250" y="410" text-anchor="middle" font-size="11" fill="#888">'
            f'Total unique: {fmt(total)}\u2002|\u2002N={size}</text>\n'
            "</svg>"
        )
        return svg

    return '<p style="color:#888;padding:12px">Not enough tools with results for Venn diagram.</p>'


def build_variant_tab(
    analysis_data: dict,   # {software: {size: gq_data_dict}}
    overlap_data:  dict,   # {size: overlap_dict}
) -> tuple[str, str]:
    """Return (html_content, js_calls) for the Variant Analysis tab."""
    html_parts: list[str] = []
    js_parts:   list[str] = []

    has_any_gq = any(
        bool(analysis_data.get(sw, {})) for sw in SOFTWARES
    )

    if not has_any_gq and not overlap_data:
        html_parts.append(
            '<p style="padding:24px;color:#888">'
            "No variant analysis data found. "
            "Run <code>scripts/03_analyze_variants.py</code> to generate."
            "</p>"
        )
        return "\n".join(html_parts), ""

    # ── GQ Histograms ─────────────────────────────────────────────────────────
    if has_any_gq:
        html_parts.append(
            '<div class="section-title">Genotype Quality (GQ) Distribution</div>'
        )
        html_parts.append(
            '<p style="color:#666;font-size:0.85rem;margin:-8px 0 12px">'
            "Each line = one dataset size (N=10 … 100). "
            "Density = fraction of genotype calls in each 5-unit GQ bin.</p>"
        )
        html_parts.append('<div class="grid-3">')

        for sw in SOFTWARES:
            div_id = f"chart_gq_{sw}"
            html_parts.append(
                f'<div class="card"><div id="{div_id}" style="height:320px"></div></div>'
            )
            sw_data = analysis_data.get(sw, {})
            size_hist_list = [
                (sz, rec["gq_histogram"])
                for sz in sorted(sw_data)
                if (rec := sw_data[sz]) and rec.get("gq_histogram")
            ]
            label = SOFT_LABEL.get(sw, sw)
            if size_hist_list:
                js_parts.append(
                    _gq_chart_js(div_id, sw, size_hist_list,
                                 f"GQ Distribution — {label}")
                )
            else:
                js_parts.append(
                    f"document.getElementById('{div_id}').innerHTML="
                    f"'<p style=\"padding:20px;color:#888\">No GQ data for {label}</p>';"
                )

        html_parts.append("</div>")  # end grid-3

    # ── Venn Diagrams ─────────────────────────────────────────────────────────
    if overlap_data:
        html_parts.append(
            '<div class="section-title">Variant Call Overlap (Venn Diagram)</div>'
        )
        available_sizes = sorted(overlap_data.keys())

        html_parts.append(
            '<div style="display:flex;flex-wrap:wrap;gap:20px;align-items:flex-start">'
        )
        for sz in available_sizes:
            od  = overlap_data[sz]
            svg = _make_venn_svg(od)
            html_parts.append(
                f'<div class="card" style="flex:0 0 auto;width:260px">'
                f'<h3 style="font-size:0.85rem;margin-bottom:6px">N = {sz}</h3>'
                f"{svg}</div>"
            )
        html_parts.append("</div>")

        # ── Overlap summary table ─────────────────────────────────────────────
        html_parts.append(
            '<div class="section-title">Variant Overlap Statistics</div>'
        )
        all_tools = sorted({
            t for od in overlap_data.values() for t in od.get("tools", [])
        })
        html_parts.append('<div class="card" style="overflow-x:auto"><table>')

        # Header
        header = ["<th>N</th>", "<th>Total unique</th>"]
        for t in all_tools:
            header.append(f"<th>{SOFT_LABEL.get(t, t)} only</th>")
        if len(all_tools) == 2:
            header.append("<th>Common</th>")
        elif len(all_tools) >= 3:
            for i in range(len(all_tools)):
                for j in range(i + 1, len(all_tools)):
                    l1 = SOFT_LABEL.get(all_tools[i], all_tools[i])
                    l2 = SOFT_LABEL.get(all_tools[j], all_tools[j])
                    header.append(f"<th>{l1} ∩ {l2} only</th>")
            header.append("<th>All common</th>")
        html_parts.append("<tr>" + "".join(header) + "</tr>")

        for sz in available_sizes:
            od       = overlap_data[sz]
            total    = od.get("total_unique", 0)
            excl     = od.get("exclusive", {})
            pairwise = od.get("pairwise_only", {})

            def _cell(val):
                if not isinstance(val, int):
                    return "<td>—</td>"
                return f"<td>{val:,} ({_pct_str(val, total)})</td>"

            row = [f"<td>{sz}</td>", f"<td>{total:,}</td>"]
            for t in all_tools:
                row.append(_cell(excl.get(t)))
            if len(all_tools) == 2:
                row.append(_cell(od.get("all_common", 0)))
            elif len(all_tools) >= 3:
                for i in range(len(all_tools)):
                    for j in range(i + 1, len(all_tools)):
                        key = f"{all_tools[i]}_{all_tools[j]}"
                        row.append(_cell(pairwise.get(key)))
                row.append(_cell(od.get("all_common", 0)))

            html_parts.append("<tr>" + "".join(row) + "</tr>")

        html_parts.append("</table></div>")

    return "\n".join(html_parts), "\n".join(js_parts)


# ─── HTML / Plotly report ─────────────────────────────────────────────────────

_CSS = """
* { box-sizing: border-box; margin: 0; padding: 0; }
body { font-family: 'Segoe UI', Arial, sans-serif; background: #f5f7fa; color: #333; }
header { background: #1a1a2e; color: #eee; padding: 24px 32px; }
header h1 { font-size: 1.6rem; font-weight: 600; }
header p  { font-size: 0.85rem; color: #aaa; margin-top: 4px; }
.tabs { display: flex; gap: 4px; padding: 16px 32px 0; background: #fff;
        border-bottom: 2px solid #e0e0e0; }
.tab-btn { padding: 8px 18px; border: none; border-radius: 6px 6px 0 0;
           cursor: pointer; background: #e8ecf0; color: #555; font-size: 0.9rem;
           transition: background 0.2s; }
.tab-btn.active { background: #1a1a2e; color: #fff; }
.tab-btn:hover:not(.active) { background: #d0d6de; }
.tab-content { display: none; padding: 24px 32px; }
.tab-content.active { display: block; }
.grid-2 { display: grid; grid-template-columns: 1fr 1fr; gap: 20px; }
.grid-3 { display: grid; grid-template-columns: 1fr 1fr 1fr; gap: 20px; }
.card  { background: #fff; border-radius: 8px; box-shadow: 0 2px 8px rgba(0,0,0,.08);
         padding: 16px; }
.card h3 { font-size: 0.95rem; color: #555; margin-bottom: 8px; }
.section-title { font-size: 1.1rem; font-weight: 600; margin: 20px 0 12px; color: #1a1a2e; }
table { width: 100%; border-collapse: collapse; font-size: 0.85rem; }
th { background: #1a1a2e; color: #fff; padding: 8px 12px; text-align: left; }
td { padding: 7px 12px; border-bottom: 1px solid #eee; }
tr:nth-child(even) td { background: #f9f9f9; }
.badge { display: inline-block; padding: 2px 8px; border-radius: 12px;
         font-size: 0.75rem; font-weight: 600; }
.badge-success { background: #d4edda; color: #155724; }
.badge-failed  { background: #f8d7da; color: #721c24; }
.badge-skipped { background: #fff3cd; color: #856404; }
.stat { text-align: center; }
.stat .val { font-size: 1.4rem; font-weight: 700; color: #1a1a2e; }
.stat .lbl { font-size: 0.75rem; color: #888; margin-top: 2px; }
@media (max-width: 900px) { .grid-2, .grid-3 { grid-template-columns: 1fr; } }
"""

_JS_TAB = """
function showTab(id, btn) {
    document.querySelectorAll('.tab-content').forEach(el => el.classList.remove('active'));
    document.querySelectorAll('.tab-btn').forEach(el => el.classList.remove('active'));
    document.getElementById(id).classList.add('active');
    btn.classList.add('active');
    // Trigger resize so Plotly reflows
    window.dispatchEvent(new Event('resize'));
}
"""

PLOTLY_CDN = "https://cdn.plot.ly/plotly-2.35.0.min.js"

_CHART_LAYOUT_BASE = {
    "margin": {"l": 60, "r": 20, "t": 40, "b": 60},
    "paper_bgcolor": "rgba(0,0,0,0)",
    "plot_bgcolor": "#fafafa",
    "font": {"family": "Segoe UI, Arial", "size": 12},
    "xaxis": {"title": "Dataset size (# GVCFs)", "gridcolor": "#eee"},
    "yaxis": {"gridcolor": "#eee"},
    "legend": {"orientation": "h", "y": -0.2},
    "hovermode": "x unified",
}


def _js_trace(sw: str, xs: list, ys: list, metric_label: str) -> str:
    return json.dumps({
        "x": xs,
        "y": [round(v, 3) for v in ys],
        "name": SOFT_LABEL.get(sw, sw),
        "mode": "lines+markers",
        "line": {"color": SOFT_COLOR.get(sw, "#888"), "width": 2},
        "marker": {"size": 7},
        "hovertemplate": f"{SOFT_LABEL.get(sw, sw)}: %{{y:.2f}}<extra></extra>",
    })


def _make_chart_js(div_id: str, traces: list[str], title: str, y_title: str) -> str:
    layout = dict(_CHART_LAYOUT_BASE)
    layout["title"] = {"text": title, "font": {"size": 14}}
    layout["yaxis"] = dict(layout.get("yaxis", {}), title=y_title, gridcolor="#eee")
    return (
        f"Plotly.newPlot('{div_id}', [{', '.join(traces)}], "
        f"{json.dumps(layout)}, {{responsive: true}});"
    )


def build_comparison_tab(data: dict) -> tuple[str, str]:
    """Return (html_content, js_calls)."""
    html_parts = ['<div class="section-title">Side-by-side comparison — all softwares</div>']
    js_parts   = []
    chart_idx  = 0

    chart_groups = [
        [("wall_time_s", "Wall Time (s)"),
         ("cpu_pct_avg", "CPU avg (%)"),
         ("cpu_pct_max", "CPU peak (%)")],
        [("ram_used_gb_avg", "RAM avg (GB)"),
         ("ram_used_gb_max", "RAM peak (GB)"),
         ("disk_used_gb_max", "Disk peak (GB)")],
        [("gpu_util_pct_avg", "GPU util avg (%)"),
         ("gpu_util_pct_max", "GPU util peak (%)"),
         ("gpu_mem_used_mb_max", "GPU Mem peak (MB)")],
    ]

    for group in chart_groups:
        html_parts.append('<div class="grid-3">')
        for metric, label in group:
            div_id = f"chart_cmp_{chart_idx}"
            chart_idx += 1
            html_parts.append(f'<div class="card"><div id="{div_id}" style="height:280px"></div></div>')
            series = build_series(data, metric)
            traces = [_js_trace(sw, xs, ys, label) for sw, (xs, ys) in series.items()]
            if traces:
                js_parts.append(_make_chart_js(div_id, traces, label, label))
            else:
                js_parts.append(f"document.getElementById('{div_id}').innerHTML='<p style=\"padding:20px;color:#888\">No data for {label}</p>';")
        html_parts.append("</div>")

    # R² summary table
    html_parts.append('<div class="section-title">Scalability Analysis (R² — linear fit)</div>')
    html_parts.append('<div class="card"><table>')
    header_cols = ["Metric"] + [SOFT_LABEL.get(s, s) for s in SOFTWARES]
    html_parts.append("<tr>" + "".join(f"<th>{c}</th>" for c in header_cols) + "</tr>")

    for metric, label in METRICS.items():
        row = [f"<td>{label}</td>"]
        for sw in SOFTWARES:
            size_dict = data.get(sw, {})
            xs = sorted(k for k in size_dict if size_dict[k].get("status") == "success")
            ys = [size_dict[s].get(metric) for s in xs]
            xs_f = [x for x, y in zip(xs, ys) if isinstance(y, (int, float))]
            ys_f = [y for y in ys if isinstance(y, (int, float))]
            r2 = compute_r2(xs_f, ys_f)
            if r2 is None:
                cell = "<td>—</td>"
            else:
                color = "#2ca02c" if r2 >= 0.95 else ("#ff7f0e" if r2 >= 0.80 else "#d62728")
                cell = f'<td style="color:{color};font-weight:600">{r2:.4f}</td>'
            row.append(cell)
        html_parts.append("<tr>" + "".join(row) + "</tr>")

    html_parts.append("</table></div>")
    return "\n".join(html_parts), "\n".join(js_parts)


def build_software_tab(sw: str, data: dict) -> tuple[str, str]:
    """HTML + JS for a single software's tab."""
    label    = SOFT_LABEL.get(sw, sw)
    size_dict = data.get(sw, {})
    sizes    = sorted(size_dict)

    html_parts = [f'<div class="section-title">{label} — Scalability Metrics</div>']
    js_parts   = []

    pairs = [
        ("wall_time_s",         "Wall Time (s)"),
        ("cpu_pct_avg",         "CPU avg (%)"),
        ("gpu_util_pct_avg",    "GPU util avg (%)"),
        ("ram_used_gb_avg",     "RAM avg (GB)"),
        ("gpu_mem_used_mb_max", "GPU Mem peak (MB)"),
        ("disk_used_gb_max",    "Disk peak (GB)"),
    ]

    html_parts.append('<div class="grid-3">')
    for metric, mlabel in pairs:
        div_id = f"chart_{sw}_{metric}"
        html_parts.append(f'<div class="card"><div id="{div_id}" style="height:260px"></div></div>')
        xs, ys = [], []
        for s in sizes:
            rec = size_dict[s]
            if rec.get("status") != "success":
                continue
            v = rec.get(metric)
            if isinstance(v, (int, float)):
                xs.append(s)
                ys.append(v)
        traces = []
        if xs:
            r2  = compute_r2(xs, ys)
            r2s = f" (R²={r2:.3f})" if r2 is not None else ""
            traces.append(_js_trace(sw, xs, ys, mlabel))
            # Add linear fit trace
            slope, intercept = _linreg(xs, ys)
            if slope is not None:
                fit_ys = [round(slope * x + intercept, 3) for x in xs]
                fit_trace = json.dumps({
                    "x": xs, "y": fit_ys,
                    "name": "Linear fit",
                    "mode": "lines",
                    "line": {"dash": "dash", "color": "#aaa", "width": 1},
                    "hoverinfo": "skip",
                })
                traces.append(fit_trace)
            js_parts.append(_make_chart_js(div_id, traces,
                                           f"{mlabel}{r2s}", mlabel))
        else:
            js_parts.append(f"document.getElementById('{div_id}').innerHTML='<p style=\"padding:20px;color:#888\">No data</p>';")
    html_parts.append("</div>")

    # Inflection points
    inf_rows = []
    for metric, mlabel in pairs:
        xs, ys = [], []
        for s in sizes:
            rec = size_dict[s]
            if rec.get("status") != "success":
                continue
            v = rec.get(metric)
            if isinstance(v, (int, float)):
                xs.append(s)
                ys.append(v)
        for inf in find_inflections(xs, ys):
            inf_rows.append({**inf, "metric": mlabel})

    if inf_rows:
        html_parts.append('<div class="section-title">Detected Inflection Points</div>')
        html_parts.append('<div class="card"><table>')
        html_parts.append("<tr><th>Metric</th><th>Dataset Size</th>"
                          "<th>Slope Before</th><th>Slope After</th></tr>")
        for row in inf_rows:
            html_parts.append(
                f"<tr><td>{row['metric']}</td><td>{row['dataset_size']}</td>"
                f"<td>{row['slope_before']:.4f}</td><td>{row['slope_after']:.4f}</td></tr>"
            )
        html_parts.append("</table></div>")

    # Detailed results table
    html_parts.append(f'<div class="section-title">{label} — Run Details</div>')
    html_parts.append('<div class="card"><table>')
    html_parts.append(
        "<tr><th>Size</th><th>Status</th><th>Wall Time (s)</th>"
        "<th>CPU avg%</th><th>GPU util%</th><th>RAM avg GB</th>"
        "<th>Variants</th></tr>"
    )
    for s in sizes:
        rec     = size_dict[s]
        status  = rec.get("status", "unknown")
        cls     = "success" if status == "success" else ("skipped" if "skip" in status else "failed")
        wt      = rec.get("wall_time_s", 0)
        cpu     = rec.get("cpu_pct_avg", "—")
        gpu     = rec.get("gpu_util_pct_avg", "—")
        ram     = rec.get("ram_used_gb_avg", "—")
        variants = rec.get("variant_count", "—")

        def _fmt(v):
            return f"{v:.2f}" if isinstance(v, float) else str(v)

        html_parts.append(
            f"<tr><td>{s}</td>"
            f'<td><span class="badge badge-{cls}">{status}</span></td>'
            f"<td>{wt:.1f}</td><td>{_fmt(cpu)}</td><td>{_fmt(gpu)}</td>"
            f"<td>{_fmt(ram)}</td><td>{variants}</td></tr>"
        )
    html_parts.append("</table></div>")

    return "\n".join(html_parts), "\n".join(js_parts)


def build_rawdata_tab(data: dict) -> str:
    rows = []
    for sw in SOFTWARES:
        size_dict = data.get(sw, {})
        for size in sorted(size_dict):
            rec = size_dict[size]
            rows.append(rec)

    html = ['<div class="section-title">Raw Benchmark Data</div>',
            '<div class="card" style="overflow-x:auto"><table>',
            "<tr><th>Software</th><th>Size</th><th>Status</th>"
            "<th>Wall Time (s)</th><th>CPU avg%</th><th>CPU peak%</th>"
            "<th>GPU util%</th><th>GPU peak%</th><th>GPU Mem MB</th>"
            "<th>RAM avg GB</th><th>RAM peak GB</th><th>Disk peak GB</th>"
            "<th>Variants</th><th>Valid</th></tr>"]

    for rec in rows:
        def g(k):
            v = rec.get(k, "—")
            return f"{v:.2f}" if isinstance(v, float) else str(v)
        sw = rec.get("software", "—")
        status = rec.get("status", "—")
        cls = "success" if status == "success" else ("skipped" if "skip" in status else "failed")
        html.append(
            f"<tr><td>{SOFT_LABEL.get(sw, sw)}</td>"
            f"<td>{rec.get('dataset_size','—')}</td>"
            f'<td><span class="badge badge-{cls}">{status}</span></td>'
            f"<td>{g('wall_time_s')}</td>"
            f"<td>{g('cpu_pct_avg')}</td><td>{g('cpu_pct_max')}</td>"
            f"<td>{g('gpu_util_pct_avg')}</td><td>{g('gpu_util_pct_max')}</td>"
            f"<td>{g('gpu_mem_used_mb_max')}</td>"
            f"<td>{g('ram_used_gb_avg')}</td><td>{g('ram_used_gb_max')}</td>"
            f"<td>{g('disk_used_gb_max')}</td>"
            f"<td>{rec.get('variant_count','—')}</td>"
            f"<td>{'✓' if rec.get('output_valid') else '✗'}</td></tr>"
        )
    html.append("</table></div>")
    return "\n".join(html)


def _render_caveats(caveats: list[str] | None) -> str:
    if not caveats:
        return ""
    items = "".join(f"<li>{c}</li>" for c in caveats)
    return (
        '<div style="background:#fff3cd;border-left:4px solid #ffc107;'
        'padding:12px 24px;margin:0 32px 16px;border-radius:4px">'
        '<strong>⚠ Data quality notes:</strong><ul style="margin:6px 0 0 16px">'
        f"{items}</ul></div>"
    )


def generate_html(
    data: dict,
    output_path: Path,
    caveats: list[str] | None = None,
    analysis_data: dict | None = None,
    overlap_data:  dict | None = None,
):
    cmp_html,   cmp_js   = build_comparison_tab(data)
    sw_tabs = {}
    sw_js   = {}
    for sw in SOFTWARES:
        sw_tabs[sw], sw_js[sw] = build_software_tab(sw, data)
    raw_html = build_rawdata_tab(data)

    var_html, var_js = build_variant_tab(
        analysis_data or {}, overlap_data or {}
    )

    # Count summary stats
    total_runs    = sum(len(sd) for sd in data.values())
    success_runs  = sum(
        sum(1 for r in sd.values() if r.get("status") == "success")
        for sd in data.values()
    )
    total_variants = sum(
        sum(r.get("variant_count", 0) or 0 for r in sd.values())
        for sd in data.values()
    )

    all_times = [
        r.get("wall_time_s", 0)
        for sd in data.values() for r in sd.values()
        if r.get("status") == "success"
    ]
    total_wall_h = round(sum(all_times) / 3600, 2) if all_times else 0

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>GVCF Scalability Benchmark Report</title>
  <script src="{PLOTLY_CDN}"></script>
  <style>{_CSS}</style>
</head>
<body>

<header>
  <h1>GVCF Scalability Benchmark Report</h1>
  <p>Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} &nbsp;|&nbsp;
     Tools: GLnexus, NVIDIA CLARA Parabricks, GATK &nbsp;|&nbsp;
     Dataset sizes: 10 – 100 GVCFs (step 10)</p>
</header>

<!-- Summary cards -->
<div style="display:flex;gap:16px;padding:20px 32px;background:#fff;border-bottom:1px solid #eee;flex-wrap:wrap">
  <div class="card stat" style="flex:1;min-width:120px">
    <div class="val">{total_runs}</div><div class="lbl">Total runs</div>
  </div>
  <div class="card stat" style="flex:1;min-width:120px">
    <div class="val">{success_runs}</div><div class="lbl">Successful</div>
  </div>
  <div class="card stat" style="flex:1;min-width:120px">
    <div class="val">{total_runs - success_runs}</div><div class="lbl">Failed / skipped</div>
  </div>
  <div class="card stat" style="flex:1;min-width:120px">
    <div class="val">{total_wall_h}h</div><div class="lbl">Total compute time</div>
  </div>
  <div class="card stat" style="flex:1;min-width:120px">
    <div class="val">{total_variants:,}</div><div class="lbl">Total variants called</div>
  </div>
</div>

{_render_caveats(caveats)}
<div class="tabs">
  <button class="tab-btn active" onclick="showTab('tab-cmp',this)">Overview</button>
  <button class="tab-btn" onclick="showTab('tab-glnexus',this)">GLnexus</button>
  <button class="tab-btn" onclick="showTab('tab-parabricks',this)">Parabricks</button>
  <button class="tab-btn" onclick="showTab('tab-gatk',this)">GATK</button>
  <button class="tab-btn" onclick="showTab('tab-variants',this)">Variant Analysis</button>
  <button class="tab-btn" onclick="showTab('tab-raw',this)">Raw Data</button>
</div>

<div id="tab-cmp" class="tab-content active">
{cmp_html}
</div>

<div id="tab-glnexus" class="tab-content">
{sw_tabs.get('glnexus','')}
</div>

<div id="tab-parabricks" class="tab-content">
{sw_tabs.get('parabricks','')}
</div>

<div id="tab-gatk" class="tab-content">
{sw_tabs.get('gatk','')}
</div>

<div id="tab-variants" class="tab-content">
{var_html}
</div>

<div id="tab-raw" class="tab-content">
{raw_html}
</div>

<footer style="padding:16px 32px;text-align:center;color:#999;font-size:0.8rem;margin-top:32px">
  GVCF Scalability Benchmark &mdash; {datetime.now().year}
</footer>

<script>
{_JS_TAB}

// ── Chart initialisation ──────────────────────────────────────────────────────
(function() {{
  // Overview tab
  {cmp_js}
  // Per-software tabs
  {sw_js.get('glnexus','')}
  {sw_js.get('parabricks','')}
  {sw_js.get('gatk','')}
  // Variant analysis tab
  {var_js}
}})();
</script>

</body>
</html>"""

    output_path.write_text(html, encoding="utf-8")


# ─── Excel report ─────────────────────────────────────────────────────────────

def _xl_header_style():
    return {
        "font":      Font(bold=True, color="FFFFFF"),
        "fill":      PatternFill("solid", fgColor="1A1A2E"),
        "alignment": Alignment(horizontal="center"),
    }


def _apply(cell, **kwargs):
    for attr, val in kwargs.items():
        setattr(cell, attr, val)


def generate_excel(data: dict, output_path: Path):
    if not HAS_OPENPYXL:
        print("[WARN] openpyxl not installed — skipping Excel export")
        print("[WARN] Install with:  pip install openpyxl")
        return

    wb = openpyxl.Workbook()

    # ── Sheet 1: Raw data ─────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Raw Data"

    columns = [
        ("Software",          "software"),
        ("Dataset Size",      "dataset_size"),
        ("Status",            "status"),
        ("Wall Time (s)",     "wall_time_s"),
        ("CPU avg %",         "cpu_pct_avg"),
        ("CPU peak %",        "cpu_pct_max"),
        ("GPU util avg %",    "gpu_util_pct_avg"),
        ("GPU util peak %",   "gpu_util_pct_max"),
        ("GPU Mem peak MB",   "gpu_mem_used_mb_max"),
        ("RAM avg GB",        "ram_used_gb_avg"),
        ("RAM peak GB",       "ram_used_gb_max"),
        ("Disk peak GB",      "disk_used_gb_max"),
        ("Variants",          "variant_count"),
        ("Output valid",      "output_valid"),
        ("Start time",        "start_time"),
        ("End time",          "end_time"),
    ]

    hs = _xl_header_style()
    for col_idx, (header, _) in enumerate(columns, start=1):
        cell = ws1.cell(row=1, column=col_idx, value=header)
        _apply(cell, font=hs["font"], fill=hs["fill"], alignment=hs["alignment"])

    row_num = 2
    for sw in SOFTWARES:
        size_dict = data.get(sw, {})
        for size in sorted(size_dict):
            rec = size_dict[size]
            for col_idx, (_, key) in enumerate(columns, start=1):
                val = rec.get(key, "")
                if isinstance(val, bool):
                    val = "Yes" if val else "No"
                ws1.cell(row=row_num, column=col_idx, value=val)
            row_num += 1

    for col in ws1.columns:
        width = max(len(str(cell.value or "")) for cell in col) + 4
        ws1.column_dimensions[col[0].column_letter].width = min(width, 30)

    # ── Sheet 2: Summary ──────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "GVCF Scalability Benchmark — Summary"
    _apply(ws2["A1"], font=Font(bold=True, size=14))
    ws2["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

    # R² table
    ws2["A4"] = "R² (linear fit) per metric"
    _apply(ws2["A4"], font=Font(bold=True))

    r2_headers = ["Metric", "GLnexus", "Parabricks", "GATK"]
    for ci, h in enumerate(r2_headers, start=1):
        cell = ws2.cell(row=5, column=ci, value=h)
        _apply(cell, font=hs["font"], fill=hs["fill"], alignment=hs["alignment"])

    row_num = 6
    for metric, label in METRICS.items():
        ws2.cell(row=row_num, column=1, value=label)
        for ci, sw in enumerate(SOFTWARES, start=2):
            size_dict = data.get(sw, {})
            xs = sorted(k for k in size_dict if size_dict[k].get("status") == "success")
            ys_raw = [size_dict[s].get(metric) for s in xs]
            xs_f = [x for x, y in zip(xs, ys_raw) if isinstance(y, (int, float))]
            ys_f = [y for y in ys_raw if isinstance(y, (int, float))]
            r2 = compute_r2(xs_f, ys_f)
            cell = ws2.cell(row=row_num, column=ci,
                            value=round(r2, 4) if r2 is not None else "N/A")
            if isinstance(r2, float):
                if r2 >= 0.95:
                    cell.fill = PatternFill("solid", fgColor="D4EDDA")
                elif r2 >= 0.80:
                    cell.fill = PatternFill("solid", fgColor="FFF3CD")
                else:
                    cell.fill = PatternFill("solid", fgColor="F8D7DA")
        row_num += 1

    # Wall-time summary table
    row_num += 2
    ws2.cell(row=row_num, column=1, value="Wall Time by Software and Size (seconds)")
    _apply(ws2.cell(row=row_num, column=1), font=Font(bold=True))
    row_num += 1

    size_set = sorted({s for sd in data.values() for s in sd})
    ws2.cell(row=row_num, column=1, value="Size \\ Software")
    for ci, sw in enumerate(SOFTWARES, start=2):
        cell = ws2.cell(row=row_num, column=ci, value=SOFT_LABEL.get(sw, sw))
        _apply(cell, font=hs["font"], fill=hs["fill"], alignment=hs["alignment"])
    row_num += 1

    for size in size_set:
        ws2.cell(row=row_num, column=1, value=size)
        for ci, sw in enumerate(SOFTWARES, start=2):
            rec = data.get(sw, {}).get(size)
            if rec and rec.get("status") == "success":
                ws2.cell(row=row_num, column=ci, value=round(rec.get("wall_time_s", 0), 1))
            else:
                ws2.cell(row=row_num, column=ci, value="N/A")
        row_num += 1

    for col in ws2.columns:
        width = max(len(str(cell.value or "")) for cell in col) + 4
        ws2.column_dimensions[col[0].column_letter].width = min(width, 30)

    wb.save(output_path)
    print(f"[REPORT] Excel written: {output_path}")


# ─── Caveats ──────────────────────────────────────────────────────────────────

def _collect_caveats(data: dict) -> list[str]:
    """
    Return human-readable warnings about data quality issues, e.g. when the
    monitoring interval (5 s) is larger than the measured wall time, meaning
    peak GPU/CPU values were almost certainly not captured.
    """
    caveats = []
    MONITOR_INTERVAL_S = 5   # matches config.sh MONITOR_INTERVAL default
    COARSE_THRESHOLD   = 3   # flag if wall_time < THRESHOLD × interval

    for sw, size_dict in data.items():
        coarse_sizes = []
        for size, rec in size_dict.items():
            if rec.get("status") != "success":
                continue
            wt = rec.get("wall_time_s", 0)
            if wt < MONITOR_INTERVAL_S * COARSE_THRESHOLD:
                coarse_sizes.append((size, wt))
        if coarse_sizes:
            sizes_str = ", ".join(
                f"N={s} ({wt:.0f}s)" for s, wt in sorted(coarse_sizes)
            )
            caveats.append(
                f"{SOFT_LABEL.get(sw, sw)}: execution faster than {COARSE_THRESHOLD}× "
                f"monitor interval ({MONITOR_INTERVAL_S}s) — GPU/CPU peaks unreliable "
                f"for {sizes_str}"
            )
    return caveats


# ─── Main ─────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Generate GVCF benchmark reports")
    parser.add_argument("--benchmark-dir", required=True,
                        help="Root benchmarks directory")
    parser.add_argument("--output-dir",    required=True,
                        help="Directory for report output")
    args = parser.parse_args()

    benchmark_dir = Path(args.benchmark_dir)
    output_dir    = Path(args.output_dir)
    metrics_dir   = benchmark_dir / "03_metrics"
    output_dir.mkdir(parents=True, exist_ok=True)

    print(f"[REPORT] Loading metrics from {metrics_dir}")
    data = load_all_metrics(metrics_dir)

    total = sum(len(sd) for sd in data.values())
    print(f"[REPORT] Loaded {total} metric records")

    analysis_data, overlap_data = load_variant_data(metrics_dir)
    n_gq  = sum(len(v) for v in analysis_data.values())
    n_ovl = len(overlap_data)
    print(f"[REPORT] Variant analysis: {n_gq} GQ histograms, {n_ovl} overlap records")

    if total == 0:
        print("[WARN] No metrics found — reports will be empty but will still be created")

    # Detect runs where the monitoring interval is too coarse relative to
    # the wall time (< 3× the interval), meaning peak resource values were
    # likely missed entirely.
    caveats = _collect_caveats(data)
    if caveats:
        print("[WARN] Coarse-monitoring caveats detected (GPU/CPU peaks may be 0):")
        for c in caveats:
            print(f"       {c}")

    html_path  = output_dir / "benchmark_report.html"
    excel_path = output_dir / "benchmark_data.xlsx"

    try:
        generate_html(data, html_path, caveats, analysis_data, overlap_data)
        print(f"[REPORT] HTML written: {html_path}")
    except Exception as exc:
        print(f"[ERROR] HTML generation failed: {exc}", file=sys.stderr)
        import traceback; traceback.print_exc()

    try:
        generate_excel(data, excel_path)
    except Exception as exc:
        print(f"[ERROR] Excel generation failed: {exc}", file=sys.stderr)
        import traceback; traceback.print_exc()

    print("[REPORT] Done.")


if __name__ == "__main__":
    main()
